import sys
import os
from PIL import Image, ImageOps, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---------------------------------------------------------
# OECO Brand Colors & Tokens
# ---------------------------------------------------------
COLOR_BG_DARK = RGBColor(0x17, 0x14, 0x12)        # #171412 Deep Espresso Brown
COLOR_CARD_BG = RGBColor(0x23, 0x1E, 0x1B)        # #231E1B Dark Cocoa Card
COLOR_CARD_BG_ALT = RGBColor(0x2A, 0x24, 0x20)    # #2A2420 Slightly lighter card
COLOR_CARD_BORDER = RGBColor(0x3D, 0x34, 0x2D)    # #3D342D Card border
COLOR_GOLD = RGBColor(0xD1, 0xB6, 0x88)           # #D1B688 OECO Champagne Gold
COLOR_GOLD_LIGHT = RGBColor(0xE8, 0xD5, 0xB5)     # #E8D5B5 Light Gold
COLOR_TEXT_LIGHT = RGBColor(0xFA, 0xF8, 0xF5)     # #FAF8F5 Off-White / Creme
COLOR_TEXT_MUTED = RGBColor(0xA8, 0x9F, 0x91)     # #A89F91 Muted Warm Gray
COLOR_TEXT_DARK = RGBColor(0x1C, 0x18, 0x15)      # #1C1815 Dark text for gold badges
COLOR_RED_ALERT = RGBColor(0xE0, 0x6A, 0x55)      # #E06A55 Coral / Alert Red
COLOR_GREEN_ACCENT = RGBColor(0x52, 0xA3, 0x74)   # #52A374 Emerald Green

# RGB Tuple for Pillow (Gold: #D1B688)
PIL_GOLD = (209, 182, 136, 255)

FONT_DISPLAY = "Playfair Display"
FONT_BODY = "Montserrat"

def crop_fit_and_round_image(image_path, target_width_in_inches, target_height_in_inches, output_name, corner_radius_px=36, centering=(0.5, 0.5), border_color=PIL_GOLD, border_width_px=4, custom_box=None):
    """
    Crops image to exact aspect ratio, applies smooth rounded corners and renders a gold border.
    """
    if not os.path.exists(image_path):
        return None
    
    cache_dir = os.path.join(os.path.dirname(image_path), "_ppt_crops")
    os.makedirs(cache_dir, exist_ok=True)
    out_path = os.path.join(cache_dir, output_name)

    target_w_px = int(target_width_in_inches * 200)
    target_h_px = int(target_height_in_inches * 200)

    with Image.open(image_path) as im:
        source_im = im
        if custom_box:
            source_im = im.crop(custom_box)

        fitted = ImageOps.fit(source_im, (target_w_px, target_h_px), method=Image.Resampling.LANCZOS, centering=centering)
        fitted_rgba = fitted.convert('RGBA')

        # Create mask for rounded corners
        mask = Image.new('L', (target_w_px, target_h_px), 0)
        draw_mask = ImageDraw.Draw(mask)
        draw_mask.rounded_rectangle(
            (0, 0, target_w_px, target_h_px),
            radius=corner_radius_px,
            fill=255
        )

        # Apply mask
        rounded_im = Image.new('RGBA', (target_w_px, target_h_px), (0, 0, 0, 0))
        rounded_im.paste(fitted_rgba, (0, 0), mask=mask)

        # Draw gold border contour
        if border_color and border_width_px > 0:
            draw_border = ImageDraw.Draw(rounded_im)
            inset = border_width_px // 2
            draw_border.rounded_rectangle(
                (inset, inset, target_w_px - inset - 1, target_h_px - inset - 1),
                radius=corner_radius_px,
                outline=border_color,
                width=border_width_px
            )

        rounded_im.save(out_path, format="PNG")
    
    return out_path

def create_presentation():
    # Base paths
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)
    images_dir = os.path.join(project_root, "public", "images")
    brand_dir = os.path.join(images_dir, "brand")

    # Raw Image asset paths
    logo_gold_path = os.path.join(brand_dir, "logo-oeco-gold.png")
    raw_strategy_path = os.path.join(images_dir, "strategy-window.jpg")
    raw_working_path = os.path.join(images_dir, "working-profile.jpg")
    raw_calm_path = os.path.join(images_dir, "calm-sofa.jpg")

    # High-quality rounded crops (without any conflicting outer shapes)
    img_s1_path = crop_fit_and_round_image(raw_strategy_path, 4.8, 5.7, "slide1_strategy_clean.png", corner_radius_px=38, centering=(0.5, 0.4))
    img_s4_path = crop_fit_and_round_image(raw_working_path, 4.5, 5.0, "slide4_working_clean.png", corner_radius_px=38, centering=(0.5, 0.45))
    img_s6_path = crop_fit_and_round_image(
        raw_calm_path, 4.0, 5.0, "slide6_calm_clean.png", 
        corner_radius_px=38, centering=(0.5, 0.45),
        custom_box=(650, 180, 1780, 2650)
    )

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Helper functions
    def set_slide_background(slide):
        bg_shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
        )
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG_DARK
        bg_shape.line.color.rgb = COLOR_BG_DARK
        return bg_shape

    def add_clean_header(slide, title_text, subtitle_text="", logo_top_right=True):
        if logo_top_right and os.path.exists(logo_gold_path):
            slide.shapes.add_picture(
                logo_gold_path, Inches(10.6), Inches(0.5), width=Inches(1.9)
            )

        # Large, mobile-readable titles
        title_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.55), Inches(9.6), Inches(1.2)
        )
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_DISPLAY
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_LIGHT

        if subtitle_text:
            p_sub = tf_t.add_paragraph()
            p_sub.text = subtitle_text
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(13.5)
            p_sub.font.color.rgb = COLOR_TEXT_MUTED
            p_sub.space_before = Pt(4)

    def add_card(slide, left, top, width, height, border_color=COLOR_CARD_BORDER, fill_color=COLOR_CARD_BG):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = fill_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # =========================================================
    # SLIDE 1: Capa & Proposta de Valor (Clean, Bold & Mobile-Optimized)
    # =========================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Logo OECO Gold (Larger for mobile view)
    if os.path.exists(logo_gold_path):
        slide1.shapes.add_picture(
            logo_gold_path, Inches(0.8), Inches(0.9), width=Inches(2.8)
        )

    # Decorative Gold Accent Line
    line = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.1), Inches(1.1), Inches(0.06)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_GOLD
    line.line.color.rgb = COLOR_GOLD

    # Main Hero Title & Subtitle (Spacious, large & cohesive)
    title_box1 = slide1.shapes.add_textbox(Inches(0.8), Inches(2.45), Inches(6.6), Inches(4.3))
    tf_hero = title_box1.text_frame
    tf_hero.word_wrap = True
    
    p_h1 = tf_hero.paragraphs[0]
    p_h1.text = "Transforme a gestão financeira da sua empresa em motor de crescimento."
    p_h1.font.name = FONT_DISPLAY
    p_h1.font.size = Pt(34)
    p_h1.font.bold = True
    p_h1.font.color.rgb = COLOR_TEXT_LIGHT
    
    p_h2 = tf_hero.add_paragraph()
    p_h2.text = "BPO Financeiro Operacional: Assumimos toda a sua rotina diária de contas a pagar, receber, conciliação e organização do caixa para você focar 100% nas vendas e nos clientes."
    p_h2.font.name = FONT_BODY
    p_h2.font.size = Pt(15)
    p_h2.font.color.rgb = COLOR_TEXT_MUTED
    p_h2.space_before = Pt(18)

    # Right Column: Clean Image with smooth rounded corners (NO floating badge)
    img_right = Inches(7.7)
    img_top = Inches(0.9)
    img_w = Inches(4.8)
    img_h = Inches(5.7)

    if img_s1_path and os.path.exists(img_s1_path):
        slide1.shapes.add_picture(img_s1_path, img_right, img_top, width=img_w, height=img_h)

    # =========================================================
    # SLIDE 2: Situação & Problema (Bullet Points & Large Mobile Fonts)
    # =========================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_clean_header(
        slide2,
        "Onde o seu tempo como empresário está sendo consumido?",
        "A maioria dos gestores gasta a maior parte da semana em tarefas operacionais em vez de liderar o crescimento."
    )

    card_data_s2 = [
        {
            "num": "01",
            "title": "Rotina Operacional Desgastante",
            "bullets": [
                "Emissão manual de boletos e notas fiscais",
                "Conciliação bancária atrasada e cheia de furos",
                "Cobrança manual e desgastante no WhatsApp"
            ],
            "highlight": "Tempo roubado do comercial e clientes"
        },
        {
            "num": "02",
            "title": "Falta de Visibilidade Real",
            "bullets": [
                "Faturamento alto, mas sem saber o lucro real",
                "Sensação de que o dinheiro 'evapora' da conta",
                "Decisões tomadas no escuro e no 'feeling'"
            ],
            "highlight": "Incerteza total sobre o saldo futuro"
        },
        {
            "num": "03",
            "title": "Alto Custo de Equipe Interna",
            "bullets": [
                "Salários, encargos trabalhistas, 13º e férias",
                "Rotatividade constante de funcionários",
                "Necessidade diária de treinamento e supervisão"
            ],
            "highlight": "Custo fixo alto e risco operacional"
        }
    ]

    card_w = Inches(3.68)
    card_h = Inches(4.9)
    card_top = Inches(1.9)

    for i, c in enumerate(card_data_s2):
        left = Inches(0.8 + i * 4.02)
        add_card(slide2, left, card_top, card_w, card_h, border_color=COLOR_CARD_BORDER)

        # Number tag
        num_box = slide2.shapes.add_textbox(left + Inches(0.25), card_top + Inches(0.2), Inches(1.0), Inches(0.4))
        p_num = num_box.text_frame.paragraphs[0]
        p_num.text = c["num"]
        p_num.font.name = FONT_DISPLAY
        p_num.font.size = Pt(24)
        p_num.font.bold = True
        p_num.font.color.rgb = COLOR_GOLD

        # Title
        t_box = slide2.shapes.add_textbox(left + Inches(0.25), card_top + Inches(0.75), card_w - Inches(0.5), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = c["title"]
        p_t.font.name = FONT_BODY
        p_t.font.size = Pt(15.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_LIGHT

        # Bullets (Large & Mobile-Friendly)
        b_box = slide2.shapes.add_textbox(left + Inches(0.25), card_top + Inches(1.6), card_w - Inches(0.5), Inches(2.2))
        tf_b = b_box.text_frame
        tf_b.word_wrap = True
        
        for idx_b, b_text in enumerate(c["bullets"]):
            p_b = tf_b.paragraphs[0] if idx_b == 0 else tf_b.add_paragraph()
            p_b.text = "• " + b_text
            p_b.font.name = FONT_BODY
            p_b.font.size = Pt(13)
            p_b.font.color.rgb = COLOR_TEXT_MUTED
            if idx_b > 0:
                p_b.space_before = Pt(8)

        # Bottom Alert Box
        hl_box = slide2.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left + Inches(0.25), card_top + Inches(3.95), card_w - Inches(0.5), Inches(0.72)
        )
        hl_box.fill.solid()
        hl_box.fill.fore_color.rgb = COLOR_CARD_BG_ALT
        hl_box.line.color.rgb = COLOR_RED_ALERT
        hl_box.line.width = Pt(1)
        tf_hl = hl_box.text_frame
        tf_hl.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_hl = tf_hl.paragraphs[0]
        p_hl.text = "⚠️ " + c["highlight"]
        p_hl.font.name = FONT_BODY
        p_hl.font.size = Pt(11)
        p_hl.font.bold = True
        p_hl.font.color.rgb = COLOR_TEXT_LIGHT
        p_hl.alignment = PP_ALIGN.CENTER

    # =========================================================
    # SLIDE 3: Implicação (I) - Mobile-Optimized
    # =========================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_clean_header(
        slide3,
        "Os riscos invisíveis de um setor financeiro desestruturado",
        "Quando o financeiro não segue um processo diário rigoroso, as perdas acontecem silenciosamente."
    )

    implications = [
        {
            "tag": "RISCO 01 · CAIXA",
            "title": "Decisões no Escuro",
            "bullets": [
                "Sem controle do saldo das próximas semanas",
                "Insegurança ao contratar ou comprar insumos",
                "O gestor descobre rombos quando já é tarde"
            ],
            "impact": "Decisões arriscadas sem base em dados reais"
        },
        {
            "tag": "RISCO 02 · PERDA FINANCEIRA",
            "title": "Multas & Inadimplência",
            "bullets": [
                "Contas pagas com atraso geram juros evitáveis",
                "Cobrança desorganizada eleva a inadimplência",
                "Vazamento contínuo da margem de lucro"
            ],
            "impact": "Perda financeira invisível todo mês"
        },
        {
            "tag": "RISCO 03 · SOBRECARGA",
            "title": "Exaustão do Gestor",
            "bullets": [
                "O dono da empresa vira o gargalo de tudo",
                "Noites e fins de semana emitindo boletos",
                "Falta de tempo para estratégia e expansão"
            ],
            "impact": "Empresa travada no limite físico do dono"
        }
    ]

    for i, imp in enumerate(implications):
        left = Inches(0.8 + i * 4.02)
        add_card(slide3, left, card_top, card_w, card_h, border_color=COLOR_RED_ALERT, fill_color=COLOR_CARD_BG)

        # Tag
        tag_box = slide3.shapes.add_textbox(left + Inches(0.25), card_top + Inches(0.2), card_w - Inches(0.5), Inches(0.35))
        p_tag = tag_box.text_frame.paragraphs[0]
        p_tag.text = imp["tag"]
        p_tag.font.name = FONT_BODY
        p_tag.font.size = Pt(10)
        p_tag.font.bold = True
        p_tag.font.color.rgb = COLOR_RED_ALERT

        # Title
        t_box = slide3.shapes.add_textbox(left + Inches(0.25), card_top + Inches(0.7), card_w - Inches(0.5), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = imp["title"]
        p_t.font.name = FONT_BODY
        p_t.font.size = Pt(15.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_TEXT_LIGHT

        # Bullets
        d_box = slide3.shapes.add_textbox(left + Inches(0.25), card_top + Inches(1.55), card_w - Inches(0.5), Inches(2.2))
        tf_d = d_box.text_frame
        tf_d.word_wrap = True
        for idx_ib, ib_text in enumerate(imp["bullets"]):
            p_ib = tf_d.paragraphs[0] if idx_ib == 0 else tf_d.add_paragraph()
            p_ib.text = "• " + ib_text
            p_ib.font.name = FONT_BODY
            p_ib.font.size = Pt(13)
            p_ib.font.color.rgb = COLOR_TEXT_MUTED
            if idx_ib > 0:
                p_ib.space_before = Pt(8)

        # Impact highlight
        imp_box = slide3.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left + Inches(0.25), card_top + Inches(3.9), card_w - Inches(0.5), Inches(0.8)
        )
        imp_box.fill.solid()
        imp_box.fill.fore_color.rgb = COLOR_CARD_BG_ALT
        imp_box.line.color.rgb = COLOR_CARD_BORDER
        imp_box.line.width = Pt(1)
        tf_imp = imp_box.text_frame
        tf_imp.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_imp = tf_imp.paragraphs[0]
        p_imp.text = "Consequência:\n" + imp["impact"]
        p_imp.font.name = FONT_BODY
        p_imp.font.size = Pt(11)
        p_imp.font.bold = True
        p_imp.font.color.rgb = COLOR_GOLD_LIGHT
        p_imp.alignment = PP_ALIGN.CENTER

    # =========================================================
    # SLIDE 4: A Solução - BPO Financeiro Operacional (Sem Card Flutuante)
    # =========================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_clean_header(
        slide4,
        "O que é o BPO Financeiro Operacional OECO?",
        "A terceirização estratégica da sua rotina diária para você focar exclusivamente em vender e crescer."
    )

    img_s4_w = Inches(4.5)
    img_s4_h = Inches(5.0)
    img_s4_left = Inches(0.8)
    img_s4_top = Inches(1.9)

    # Clean Image with smooth rounded corners (NO floating badge)
    if img_s4_path and os.path.exists(img_s4_path):
        slide4.shapes.add_picture(img_s4_path, img_s4_left, img_s4_top, width=img_s4_w, height=img_s4_h)

    # Right Column: 3 Pillars of Operational BPO (Large, readable)
    right_s4_left = Inches(5.6)
    right_s4_w = Inches(6.933)

    pillars = [
        ("1. Rotina Diária Padronizada & Sem Atalhos", "Assumimos toda a operação burocrática: agendamento de pagamentos no banco, conciliação e emissão de notas com disciplina diária."),
        ("2. Sinergia Total com a sua Contabilidade", "Não substituímos o seu contador: trabalhamos em parceria, entregando extratos 100% conciliados e notas fiscais organizadas todo mês."),
        ("3. Transparência & Relatórios de Caixa", "Relatórios periódicos claros de movimentação e posição de contas para você ter total tranquilidade do caixa sem fazer contas manuais.")
    ]

    for j, (p_title, p_desc) in enumerate(pillars):
        p_top = card_top + Inches(j * 1.7)
        add_card(slide4, right_s4_left, p_top, right_s4_w, Inches(1.55), border_color=COLOR_CARD_BORDER)
        
        t_box = slide4.shapes.add_textbox(right_s4_left + Inches(0.3), p_top + Inches(0.18), right_s4_w - Inches(0.6), Inches(1.2))
        tf_p = t_box.text_frame
        tf_p.word_wrap = True
        
        p1 = tf_p.paragraphs[0]
        p1.text = p_title
        p1.font.name = FONT_BODY
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_GOLD
        
        p2 = tf_p.add_paragraph()
        p2.text = p_desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(5)

    # =========================================================
    # SLIDE 5: Nossos Serviços no Dia a Dia (Escopo Operacional)
    # =========================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_clean_header(
        slide5,
        "O que fazemos pela sua empresa no dia a dia",
        "Uma esteira financeira operacional completa, organizada e executada com precisão diária."
    )

    services = [
        {
            "icon": "💳",
            "title": "Contas a Pagar & Receber",
            "items": [
                "Agendamento de pagamentos no banco",
                "Controle rigoroso de vencimentos",
                "Emissão de boletos e notas fiscais",
                "Baixa automática e conferência"
            ]
        },
        {
            "icon": "🏦",
            "title": "Conciliação Bancária Diária",
            "items": [
                "Conferência diária saldo x extrato",
                "Eliminação de divergências e furos",
                "Classificação no plano de contas",
                "Organização de comprovantes"
            ]
        },
        {
            "icon": "📞",
            "title": "Gestão de Inadimplência",
            "items": [
                "Régua preventiva e amigável",
                "Avisos automáticos de vencimento",
                "Cobrança ativa de títulos atrasados",
                "Recuperação rápida de liquidez"
            ]
        },
        {
            "icon": "📊",
            "title": "Organização & Relatórios",
            "items": [
                "Relatório de Contas Pagas e Recebidas",
                "Posição de Saldo e Fluxo Diário",
                "Documentos limpos p/ contabilidade",
                "Segurança e sigilo total dos dados"
            ]
        }
    ]

    s_card_w = Inches(2.78)
    s_card_h = Inches(4.9)

    for i, s in enumerate(services):
        left = Inches(0.8 + i * 2.98)
        add_card(slide5, left, card_top, s_card_w, s_card_h, border_color=COLOR_CARD_BORDER)

        # Header of service card
        sh_box = slide5.shapes.add_textbox(left + Inches(0.2), card_top + Inches(0.25), s_card_w - Inches(0.4), Inches(0.9))
        tf_sh = sh_box.text_frame
        tf_sh.word_wrap = True
        
        p_icon = tf_sh.paragraphs[0]
        p_icon.text = s["icon"] + " " + s["title"]
        p_icon.font.name = FONT_BODY
        p_icon.font.size = Pt(14)
        p_icon.font.bold = True
        p_icon.font.color.rgb = COLOR_GOLD_LIGHT

        # Divider line
        c_line = slide5.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left + Inches(0.2), card_top + Inches(1.25), s_card_w - Inches(0.4), Inches(0.02)
        )
        c_line.fill.solid()
        c_line.fill.fore_color.rgb = COLOR_CARD_BORDER
        c_line.line.color.rgb = COLOR_CARD_BORDER

        # Items
        items_box = slide5.shapes.add_textbox(left + Inches(0.2), card_top + Inches(1.4), s_card_w - Inches(0.4), Inches(3.2))
        tf_it = items_box.text_frame
        tf_it.word_wrap = True
        
        for idx, item in enumerate(s["items"]):
            p_it = tf_it.paragraphs[0] if idx == 0 else tf_it.add_paragraph()
            p_it.text = "• " + item
            p_it.font.name = FONT_BODY
            p_it.font.size = Pt(12)
            p_it.font.color.rgb = COLOR_TEXT_LIGHT
            if idx > 0:
                p_it.space_before = Pt(10)

    # =========================================================
    # SLIDE 6: Resultados & Benefícios (Sem Card Flutuante na Foto)
    # =========================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_clean_header(
        slide6,
        "O que a sua empresa ganha na prática?",
        "Ganhos operacionais, economia e tranquilidade perceptíveis desde o primeiro mês."
    )

    # Left: 4 Metric Cards in 2x2 Grid (100% Operacional)
    grid_left = Inches(0.8)
    grid_top = Inches(1.9)
    m_w = Inches(3.7)
    m_h = Inches(2.35)

    metrics = [
        {
            "badge": "MAIS TEMPO",
            "stat": "+20h",
            "stat_label": "Livres por mês",
            "desc": "O gestor deixa o operacional financeiro e volta a focar em vendas e clientes.",
            "color": COLOR_GOLD
        },
        {
            "badge": "ECONOMIA",
            "stat": "Até 60%",
            "stat_label": "Redução de custos",
            "desc": "Eliminação de encargos, férias, 13º, rescisões e softwares duplicados.",
            "color": COLOR_GREEN_ACCENT
        },
        {
            "badge": "CONCILIAÇÃO",
            "stat": "100%",
            "stat_label": "Conciliado & Seguro",
            "desc": "Conferência diária de cada centavo, eliminando divergências e pontas soltas.",
            "color": COLOR_GOLD_LIGHT
        },
        {
            "badge": "PONTUALIDADE",
            "stat": "0 Multas",
            "stat_label": "Zero juros por atraso",
            "desc": "Contas agendadas no prazo rigoroso e cobrança preventiva ativa de clientes.",
            "color": COLOR_GOLD
        }
    ]

    for idx, m in enumerate(metrics):
        col = idx % 2
        row = idx // 2
        c_left = grid_left + Inches(col * 3.9)
        c_top = grid_top + Inches(row * 2.5)

        add_card(slide6, c_left, c_top, m_w, m_h, border_color=COLOR_CARD_BORDER)

        # Stat & Badge
        b_box = slide6.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.15), m_w - Inches(0.4), Inches(0.3))
        p_b = b_box.text_frame.paragraphs[0]
        p_b.text = m["badge"]
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = m["color"]

        s_box = slide6.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.45), m_w - Inches(0.4), Inches(0.65))
        p_s = s_box.text_frame.paragraphs[0]
        p_s.text = m["stat"]
        p_s.font.name = FONT_DISPLAY
        p_s.font.size = Pt(29)
        p_s.font.bold = True
        p_s.font.color.rgb = COLOR_TEXT_LIGHT

        sl_box = slide6.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(1.15), m_w - Inches(0.4), Inches(0.3))
        p_sl = sl_box.text_frame.paragraphs[0]
        p_sl.text = m["stat_label"]
        p_sl.font.name = FONT_BODY
        p_sl.font.size = Pt(12)
        p_sl.font.bold = True
        p_sl.font.color.rgb = COLOR_GOLD_LIGHT

        d_box = slide6.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(1.5), m_w - Inches(0.4), Inches(0.75))
        tf_d = d_box.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = m["desc"]
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(11.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED

    # Right: Photo Frame (Calm Sofa - Perfect Clean Image, NO floating badge)
    photo_s6_left = Inches(8.8)
    photo_s6_top = Inches(1.9)
    photo_s6_w = Inches(3.733)
    photo_s6_h = Inches(4.9)

    if img_s6_path and os.path.exists(img_s6_path):
        slide6.shapes.add_picture(img_s6_path, photo_s6_left, photo_s6_top, width=photo_s6_w, height=photo_s6_h)

    # =========================================================
    # SLIDE 7: Próximo Passo & CTA ÚNICO (Laser-Focused Single Goal)
    # =========================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_clean_header(
        slide7,
        "Pronto para blindar a rotina financeira da sua empresa?",
        "Dê o primeiro passo em 2 minutos para receber um pré-diagnóstico personalizado diretamente no seu WhatsApp."
    )

    # Single Centered Executive Hero Card
    hero_card_w = Inches(11.733)
    hero_card_h = Inches(4.9)
    hero_card_left = Inches(0.8)
    hero_card_top = Inches(1.9)

    add_card(slide7, hero_card_left, hero_card_top, hero_card_w, hero_card_h, border_color=COLOR_GOLD, fill_color=COLOR_CARD_BG_ALT)

    # Content inside Unified Card
    tf_h = slide7.shapes.add_textbox(hero_card_left + Inches(0.6), hero_card_top + Inches(0.35), hero_card_w - Inches(1.2), Inches(4.2)).text_frame
    tf_h.word_wrap = True

    # Badge
    p_hb = tf_h.paragraphs[0]
    p_hb.text = "PASSO ÚNICO · 100% ONLINE E GRATUITO"
    p_hb.font.name = FONT_BODY
    p_hb.font.size = Pt(10.5)
    p_hb.font.bold = True
    p_hb.font.color.rgb = COLOR_GOLD
    p_hb.alignment = PP_ALIGN.CENTER

    # Title
    p_ht = tf_h.add_paragraph()
    p_ht.text = "Faça o Pré-Diagnóstico de Caixa & Rotina da sua Empresa"
    p_ht.font.name = FONT_DISPLAY
    p_ht.font.size = Pt(25)
    p_ht.font.bold = True
    p_ht.font.color.rgb = COLOR_TEXT_LIGHT
    p_ht.space_before = Pt(6)
    p_ht.alignment = PP_ALIGN.CENTER

    # Subtitle
    p_hd = tf_h.add_paragraph()
    p_hd.text = "Responda a 4 perguntas rápidas para identificarmos seus maiores gargalos de tempo e desenharmos o plano ideal para assumir sua rotina financeira."
    p_hd.font.name = FONT_BODY
    p_hd.font.size = Pt(13)
    p_hd.font.color.rgb = COLOR_TEXT_MUTED
    p_hd.space_before = Pt(8)
    p_hd.alignment = PP_ALIGN.CENTER

    # 3 Highlights in 1 Row
    p_hl = tf_h.add_paragraph()
    p_hl.text = "⚡ Menos de 2 minutos  ·  📊 Devolutiva Técnica Personalizada  ·  🔒 Sigilo e Confidencialidade Total"
    p_hl.font.name = FONT_BODY
    p_hl.font.size = Pt(12)
    p_hl.font.bold = True
    p_hl.font.color.rgb = COLOR_GOLD_LIGHT
    p_hl.space_before = Pt(14)
    p_hl.alignment = PP_ALIGN.CENTER

    # Super Big Gold Action Button (Centered)
    btn_w = Inches(8.0)
    btn_h = Inches(0.8)
    btn_left = hero_card_left + (hero_card_w - btn_w) / 2
    btn_top = hero_card_top + Inches(2.8)

    btn = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, btn_left, btn_top, btn_w, btn_h
    )
    btn.fill.solid()
    btn.fill.fore_color.rgb = COLOR_GOLD
    btn.line.color.rgb = COLOR_GOLD_LIGHT
    btn.line.width = Pt(1)
    btn.click_action.hyperlink.address = "https://oecofs.com/diagnostico"

    tf_btn = btn.text_frame
    tf_btn.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_btn = tf_btn.paragraphs[0]
    p_btn.text = "👉 CLIQUE AQUI PARA INICIAR O PRÉ-DIAGNÓSTICO"
    p_btn.font.name = FONT_BODY
    p_btn.font.size = Pt(14)
    p_btn.font.bold = True
    p_btn.font.color.rgb = COLOR_TEXT_DARK
    p_btn.alignment = PP_ALIGN.CENTER

    # Discrete Footer Information inside card
    foot_box = slide7.shapes.add_textbox(hero_card_left + Inches(0.6), hero_card_top + Inches(3.8), hero_card_w - Inches(1.2), Inches(0.8))
    tf_f = foot_box.text_frame
    tf_f.word_wrap = True

    p_f1 = tf_f.paragraphs[0]
    p_f1.text = "Ao concluir o formulário, suas respostas serão enviadas diretamente pelo WhatsApp para alinharmos os detalhes."
    p_f1.font.name = FONT_BODY
    p_f1.font.size = Pt(10.5)
    p_f1.font.color.rgb = COLOR_TEXT_MUTED
    p_f1.alignment = PP_ALIGN.CENTER

    p_f2 = tf_f.add_paragraph()
    p_f2.text = "OECO Financial Solutions · oecofs.com · contato@oecofs.com · (21) 97341-3967"
    p_f2.font.name = FONT_BODY
    p_f2.font.size = Pt(10)
    p_f2.font.color.rgb = COLOR_GOLD
    p_f2.space_before = Pt(3)
    p_f2.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_filename = "Apresentacao_BPO_Financeiro_OECO.pptx"
    output_path = os.path.join(project_root, output_filename)
    try:
        prs.save(output_path)
        print(f"Apresentação BPO Operacional gerada com sucesso em: {output_path}")
    except PermissionError:
        output_filename_v2 = "Apresentacao_BPO_Financeiro_OECO_v2.pptx"
        output_path_v2 = os.path.join(project_root, output_filename_v2)
        prs.save(output_path_v2)
        print(f"O arquivo original estava aberto. Nova versão salva em: {output_path_v2}")

if __name__ == "__main__":
    create_presentation()
